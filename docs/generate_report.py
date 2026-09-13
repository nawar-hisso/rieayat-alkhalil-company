# -*- coding: utf-8 -*-
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import os

OUT_PATH = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "بيانات-مطلوبة-لموقع-شركة-رعاية-الخليل.xlsx",
)

# Brand palette (derived from the supplied Alkhalil logo artwork)
CHAMPAGNE = "B08D5C"
CHAMPAGNE_DEEP = "8A6D42"
PEARL = "F8F6F2"
PEARL_2 = "F1EFE9"
BEIGE = "EAE3D6"
CHARCOAL = "262420"
WHITE = "FFFFFF"

header_font = Font(name="Calibri", size=11, bold=True, color=WHITE)
header_fill = PatternFill("solid", fgColor=CHAMPAGNE_DEEP)
title_font = Font(name="Calibri", size=16, bold=True, color=CHAMPAGNE_DEEP)
subtitle_font = Font(name="Calibri", size=10, italic=True, color="6F6A62")
body_font = Font(name="Calibri", size=11, color=CHARCOAL)
bold_body_font = Font(name="Calibri", size=11, bold=True, color=CHARCOAL)
note_font = Font(name="Calibri", size=10, italic=True, color="6F6A62")

fill_even = PatternFill("solid", fgColor=PEARL)
fill_odd = PatternFill("solid", fgColor=PEARL_2)
fill_priority_required = PatternFill("solid", fgColor="F3D9CB")
fill_priority_important = PatternFill("solid", fgColor="F6E7C9")
fill_priority_optional = PatternFill("solid", fgColor="E4EFE1")
fill_priority_confirm = PatternFill("solid", fgColor=BEIGE)

thin = Side(style="thin", color="D8CDB8")
border_all = Border(left=thin, right=thin, top=thin, bottom=thin)

wrap_top = Alignment(horizontal="right", vertical="top", wrap_text=True)
center = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_header_row(ws: Worksheet, row: int, ncols: int):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border_all


def style_data_row(ws: Worksheet, row: int, ncols: int, zebra_fill):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = body_font
        cell.alignment = wrap_top
        cell.fill = zebra_fill
        cell.border = border_all


def set_widths(ws: Worksheet, widths: dict):
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w


wb = Workbook()
wb.properties.title = "بيانات مطلوبة لموقع شركة رعاية الخليل"
wb.properties.creator = "Rieayat Alkhalil Company Website Project"

# ---------------------------------------------------------------
# Sheet 1: المعلومات المؤكدة
# ---------------------------------------------------------------
ws1 = wb.active
ws1.title = "المعلومات المؤكدة"
ws1.sheet_view.rightToLeft = True

ws1.merge_cells("A1:D1")
ws1["A1"] = "المعلومات المؤكدة عن شركة رعاية الخليل"
ws1["A1"].font = title_font
ws1["A1"].alignment = Alignment(horizontal="right", vertical="center")
ws1.row_dimensions[1].height = 28

ws1.merge_cells("A2:D2")
ws1["A2"] = "هذه المعلومات مؤكدة من صفحة فيسبوك الرسمية للشركة أو مزوَّدة مباشرة من المالك — لا حاجة لتكرار طلبها."
ws1["A2"].font = subtitle_font
ws1["A2"].alignment = Alignment(horizontal="right", vertical="center")
ws1.row_dimensions[2].height = 20

headers1 = ["الفئة", "المعلومة", "القيمة", "المصدر / الملاحظة"]
for i, h in enumerate(headers1, start=1):
    ws1.cell(row=3, column=i, value=h)
style_header_row(ws1, 3, 4)
ws1.freeze_panes = "A4"
ws1.auto_filter.ref = "A3:D3"

rows1 = [
    ("الشركة", "اسم الشركة بالعربية", "شركة رعاية الخليل", "مؤكد من المالك"),
    ("الشركة", "الاسم المستخدم بالإنجليزية", "Rieayat Alkhalil Company", "مستخدم في اتصالات الشركة"),
    ("الشركة", "اسم صفحة فيسبوك", "Rieayat Alkhalil Company Erbil", "facebook.com"),
    ("الشركة", "الاسم الظاهر على الشعار", "ALKHALIL COMPANY", "من ملفات الشعار المرسلة"),
    ("النشاط", "النشاط التجاري", "استيراد وبيع المواد الأولية البلاستيكية", "من بيانات فيسبوك"),
    ("النشاط", "تصنيف النشاط على فيسبوك", "تجاري وصناعي", "من بيانات فيسبوك"),
    ("الموقع الجغرافي", "المدينة والدولة", "أربيل، محافظة أربيل، العراق", "مؤكد"),
    ("الموقع الجغرافي", "إشارة خرائط جوجل", "5XW7+GHP, Erbil, Erbil Governorate, Iraq", "من رابط خرائط جوجل"),
    ("التواصل", "الهاتف الرئيسي", "+964 750 500 0335", "الرقم المعلن على فيسبوك"),
    ("التواصل", "الهاتف الثاني", "+964 750 488 8836", "مزوَّد من المالك"),
    ("التواصل", "البريد الإلكتروني", "alkhalilerbil@gmail.com", "مؤكد"),
    ("التواصل", "رابط فيسبوك", "https://www.facebook.com/people/Rieayat-Alkhalil-Company-Erbil/100083525254385/", "مؤكد"),
    ("التواصل", "رابط خرائط جوجل", "https://maps.app.goo.gl/CW4PprFEVnpu5nXb7?g_st=iwb", "مؤكد"),
    ("التوريد", "دول التوريد المعروفة", "إيران، السعودية، تركيا، الصين", "مؤكد من المالك"),
    ("الدوام", "ساعات الدوام المعلنة حالياً", "السبت–الخميس تقريباً 10:20 صباحاً–6:00 مساءً، الجمعة مغلق", "معلومة عامة من فيسبوك — يفضّل تأكيدها (انظر الورقة الثانية)"),
]

r = 4
for cat, info, val, note in rows1:
    ws1.cell(row=r, column=1, value=cat)
    ws1.cell(row=r, column=2, value=info)
    ws1.cell(row=r, column=3, value=val)
    ws1.cell(row=r, column=4, value=note)
    zebra = fill_even if r % 2 == 0 else fill_odd
    style_data_row(ws1, r, 4, zebra)
    ws1.cell(row=r, column=1).font = bold_body_font
    r += 1

set_widths(ws1, {1: 16, 2: 26, 3: 46, 4: 36})
for row in range(4, r):
    ws1.row_dimensions[row].height = 32

# ---------------------------------------------------------------
# Sheet 2: المعلومات المطلوبة
# ---------------------------------------------------------------
ws2 = wb.create_sheet("المعلومات المطلوبة")
ws2.sheet_view.rightToLeft = True

ws2.merge_cells("A1:H1")
ws2["A1"] = "المعلومات المطلوبة من مالك الشركة"
ws2["A1"].font = title_font
ws2["A1"].alignment = Alignment(horizontal="right", vertical="center")
ws2.row_dimensions[1].height = 28

ws2.merge_cells("A2:H2")
ws2["A2"] = (
    "يرجى تعبئة عمود «الإجابة» لكل نقطة. هذه الأسئلة فقط لما هو غير مؤكد بعد — "
    "الموقع يعمل حالياً بشكل كامل بالاعتماد على المعلومات المتوفرة."
)
ws2["A2"].font = subtitle_font
ws2["A2"].alignment = Alignment(horizontal="right", vertical="center")
ws2.row_dimensions[2].height = 20

headers2 = [
    "رقم", "الفئة", "السؤال", "لماذا نحتاج هذه المعلومة؟",
    "أين ستُستخدم في الموقع؟", "الأولوية", "الإجابة", "ملاحظات",
]
for i, h in enumerate(headers2, start=1):
    ws2.cell(row=3, column=i, value=h)
style_header_row(ws2, 3, 8)
ws2.freeze_panes = "A4"
ws2.auto_filter.ref = "A3:H3"

items2 = [
    (
        "الاسم والهوية",
        "هل الاسم الذي يجب اعتماده بشكل أساسي في الموقع هو \"Rieayat Alkhalil Company\" أم \"ALKHALIL COMPANY\" كما هو ظاهر في الشعار؟",
        "لضمان استخدام الاسم الصحيح في جميع عناوين ونصوص الموقع",
        "العنوان الرئيسي، الشعار، بيانات تحسين محركات البحث (SEO)",
        "ضروري",
        "يمكن اعتماد أحدهما كاسم أساسي والآخر كاسم فرعي إذا رغبتم.",
    ),
    (
        "المنتجات",
        "ما هي المواد البلاستيكية التي تبيعها الشركة فعلياً؟ يرجى ذكر الأسماء والأنواع أو الدرجات المهمة إن وجدت.",
        "الموقع الحالي يعرض فئات شائعة لهذا النشاط، وهي تقديرية وليست قائمة مؤكدة من الشركة",
        "قسم «المواد» في الصفحة الرئيسية",
        "ضروري",
        "مثال: PP، HDPE، LDPE، LLDPE، PVC، PET، Masterbatch. يُذكر فقط ما تتعامل به الشركة فعلياً.",
    ),
    (
        "أسلوب البيع",
        "هل يتم البيع بالجملة فقط، أم توجد أنماط بيع أخرى (كميات صغيرة، عقود مستمرة، إلخ)؟",
        "لتوضيح طريقة التعامل المتوقعة مع العملاء المحتملين",
        "قسم «المواد» وقسم «خطوات طلب عرض السعر»",
        "مهم",
        "",
    ),
    (
        "العملاء",
        "ما هي القطاعات أو أنواع العملاء الذين تتعامل معهم الشركة فعلياً؟",
        "الموقع الحالي يعرض تطبيقات شائعة للمواد وليست قطاعات عملاء مؤكدة بالاسم",
        "قسم «القطاعات»",
        "مهم",
        "مثال: مصانع التغليف، الحقن البلاستيكي، الأنابيب، العبوات، التجار والموزعون.",
    ),
    (
        "نطاق التغطية",
        "ما هي المناطق التي تصل إليها خدمات الشركة؟",
        "لتحديد نطاق الخدمة الجغرافي الحقيقي بدقة في نصوص الموقع",
        "قسم «من نحن» وقسم «تواصل معنا»",
        "مهم",
        "مثال: أربيل فقط، إقليم كردستان، جميع العراق.",
    ),
    (
        "التوصيل",
        "هل توجد خدمة توصيل للمواد؟ وإذا كانت متوفرة، ضمن أي نطاق؟",
        "لإضافة معلومة دقيقة حول التوصيل بدون افتراضها",
        "قسم «من نحن»",
        "اختياري",
        "",
    ),
    (
        "المخزون",
        "هل يوجد مستودع أو مخزون محلي دائم في أربيل، أم يتم التوريد حسب الطلب لكل عملية استيراد؟",
        "لتحديد الصياغة الصحيحة حول توفر المواد وتجنّب الوعد بمخزون دائم غير مؤكد",
        "قسم «من نحن» وقسم «المواد»",
        "مهم",
        "",
    ),
    (
        "طلبات خاصة",
        "هل يمكن تأمين مواد أو درجات أو مناشئ حسب طلب العميل تحديداً، خارج القائمة المعروضة؟",
        "لتأكيد رسالة «أرسل طلبك الخاص» الظاهرة حالياً في الموقع",
        "قسم «المواد» وقسم «طلب عرض سعر»",
        "اختياري",
        "",
    ),
    (
        "الموردون والعلامات التجارية",
        "هل توجد أسماء مصانع أو علامات تجارية موردة يمكن ذكرها علناً في الموقع؟",
        "لا يمكن ذكر أي جهة موردة في الموقع بدون إذن واضح من الشركة",
        "قسم «مصادر التوريد»",
        "اختياري",
        "لن يُذكر أي اسم مورّد إلا بعد تأكيد واضح بالسماح بذلك.",
    ),
    (
        "وكالات واعتمادات",
        "هل تحمل الشركة وكالة رسمية أو توزيعاً معتمداً من أي جهة؟ هل توجد شهادات يمكن نشرها؟",
        "لم تُذكر أي وكالة أو شهادة في الموقع حتى الآن لتجنّب أي معلومة غير مؤكدة",
        "قسم «لماذا رعاية الخليل»",
        "اختياري",
        "تُذكر فقط إذا كانت مؤكدة ويمكن إثباتها.",
    ),
    (
        "تاريخ التأسيس",
        "في أي سنة تأسست الشركة؟ وهل ترغبون بإظهار سنة التأسيس أو مدة الخبرة في الموقع؟",
        "لم تُستخدم أي سنة تأسيس أو مدة خبرة حالياً لتجنّب الافتراض",
        "قسم «من نحن»",
        "اختياري",
        "",
    ),
    (
        "العنوان التفصيلي",
        "ما هو العنوان التفصيلي الكامل (الحي، الشارع، أقرب معلم) إضافةً إلى رابط خرائط جوجل الحالي؟",
        "لتحسين وضوح الموقع الجغرافي للزوار والعملاء",
        "قسم «تواصل معنا»",
        "مهم",
        "",
    ),
    (
        "تأكيد ساعات الدوام",
        "هل ساعات الدوام المعلنة حالياً على فيسبوك (السبت–الخميس، تقريباً 10:20 صباحاً–6:00 مساءً، الجمعة مغلق) لا تزال صحيحة؟",
        "معلومة عامة متوفرة لكن يفضّل تأكيدها لضمان دقتها في الموقع",
        "قسم «تواصل معنا»",
        "تأكيد فقط",
        "",
    ),
    (
        "رقم واتساب الرسمي",
        "أي الرقمين هو رقم واتساب الرسمي المستخدم حالياً في زر «اطلب عرض سعر»؟ وهل الرقم الثاني مفعّل على واتساب أيضاً؟",
        "الموقع يستخدم حالياً +964 750 500 0335 فقط لجميع طلبات عرض السعر عبر واتساب",
        "زر «اطلب عرض سعر» في جميع صفحات الموقع",
        "ضروري",
        "الموقع حالياً يوجّه كل طلبات واتساب إلى +964 750 500 0335.",
    ),
    (
        "وسيلة التواصل المفضلة",
        "ما هي وسيلة التواصل التي يفضّل المالك التركيز عليها بشكل أساسي: واتساب أم الاتصال المباشر؟",
        "لضبط ترتيب وتأكيد الأزرار الرئيسية في الموقع",
        "الرأس، الشريط السفلي في الجوال، قسم التواصل",
        "اختياري",
        "",
    ),
    (
        "صور حقيقية",
        "هل تتوفر صور حقيقية للمستودع أو المواد أو الشحنات يمكن استخدامها في الموقع؟",
        "الصور الحالية في الموقع هي صور الهوية/الشعار المرسلة فقط",
        "قسم «من نحن» والصفحة الرئيسية",
        "اختياري",
        "صور بجودة جيدة تحسّن مصداقية الموقع بشكل كبير.",
    ),
    (
        "مراجعة أخيرة",
        "هل توجد أي معلومة معروضة حالياً في الموقع لا ترغب الشركة بنشرها علناً؟",
        "لضمان مراجعة أخيرة قبل الإطلاق النهائي للموقع",
        "جميع أقسام الموقع",
        "مهم",
        "",
    ),
]

priority_fill_map = {
    "ضروري": fill_priority_required,
    "مهم": fill_priority_important,
    "اختياري": fill_priority_optional,
    "تأكيد فقط": fill_priority_confirm,
}

r = 4
for idx, (cat, question, why, where, priority, note) in enumerate(items2, start=1):
    ws2.cell(row=r, column=1, value=idx)
    ws2.cell(row=r, column=2, value=cat)
    ws2.cell(row=r, column=3, value=question)
    ws2.cell(row=r, column=4, value=why)
    ws2.cell(row=r, column=5, value=where)
    ws2.cell(row=r, column=6, value=priority)
    ws2.cell(row=r, column=7, value="")
    ws2.cell(row=r, column=8, value=note)

    zebra = fill_even if r % 2 == 0 else fill_odd
    style_data_row(ws2, r, 8, zebra)
    ws2.cell(row=r, column=1).alignment = center
    ws2.cell(row=r, column=2).font = bold_body_font
    ws2.cell(row=r, column=6).font = bold_body_font
    ws2.cell(row=r, column=6).alignment = center
    ws2.cell(row=r, column=6).fill = priority_fill_map.get(priority, zebra)
    ws2.cell(row=r, column=8).font = note_font
    r += 1

set_widths(ws2, {1: 6, 2: 20, 3: 46, 4: 34, 5: 26, 6: 12, 7: 26, 8: 34})
for row in range(4, r):
    ws2.row_dimensions[row].height = 62

# ---------------------------------------------------------------
# Sheet 3: ملاحظات للموقع
# ---------------------------------------------------------------
ws3 = wb.create_sheet("ملاحظات للموقع")
ws3.sheet_view.rightToLeft = True

ws3.merge_cells("A1:B1")
ws3["A1"] = "ملاحظات عامة حول الموقع الحالي"
ws3["A1"].font = title_font
ws3["A1"].alignment = Alignment(horizontal="right", vertical="center")
ws3.row_dimensions[1].height = 28

headers3 = ["#", "ملاحظة"]
for i, h in enumerate(headers3, start=1):
    ws3.cell(row=2, column=i, value=h)
style_header_row(ws3, 2, 2)
ws3.freeze_panes = "A3"

notes3 = [
    "تم بناء الموقع الحالي بالاعتماد على المعلومات المتوفرة في صفحة فيسبوك الرسمية للشركة والمعلومات التي زوَّدنا بها المالك مباشرة.",
    "تم إنشاء بعض النصوص التسويقية (مثل قسم «من نحن» ووصف خطوات التعامل) بشكل مهني بناءً على طبيعة نشاط الشركة، دون اختلاق معلومات غير مؤكدة.",
    "لم يتم ذكر أي شهادات، أرقام إحصائية، عملاء بالاسم، وكلاء معتمدين، أو شراكات لم تُؤكَّد من الشركة.",
    "تعبئة ورقة «المعلومات المطلوبة» لاحقاً ستسمح بتحسين دقة محتوى الموقع بشكل ملموس، دون الحاجة لإعادة بناء الموقع.",
    "قائمة المواد البلاستيكية المعروضة حالياً في الموقع (البولي إيثيلين، البولي بروبيلين، PVC، PET، الماستر باتش والإضافات) تمثل فئات شائعة في هذا النوع من النشاط التجاري، ويجب تأكيدها أو تعديلها من الشركة.",
    "الصور المستخدمة حالياً في الموقع هي صور شعار وهوية الشركة المرسلة من المالك، ويمكن تحسين الموقع مستقبلاً بإضافة صور حقيقية للمستودع أو المواد أو عمليات الشحن.",
    "أرقام الهاتف وساعات الدوام مركَّزة في ملف إعدادات واحد داخل كود الموقع، مما يسهّل تحديثها لاحقاً دون الحاجة لإعادة بناء الموقع بالكامل.",
]

r = 3
for i, note in enumerate(notes3, start=1):
    ws3.cell(row=r, column=1, value=i)
    ws3.cell(row=r, column=2, value=note)
    zebra = fill_even if r % 2 == 0 else fill_odd
    style_data_row(ws3, r, 2, zebra)
    ws3.cell(row=r, column=1).alignment = center
    ws3.row_dimensions[r].height = 44
    r += 1

set_widths(ws3, {1: 6, 2: 100})

wb.save(OUT_PATH)
print("saved:", OUT_PATH)
